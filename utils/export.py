import os
from datetime import datetime

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def _set_normal_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(34, 34, 34)


def _add_meta_line(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    paragraph.add_run(value)


def generate_word_minutes(transcript_data: dict, analysis) -> str:
    """
    Generate a readable Word document with the meeting summary and minutes.
    Returns the full path of the saved .docx file.
    """
    if not analysis or not transcript_data:
        st.error("No analysis data available for export")
        return ""

    with st.spinner("Generating Word meeting summary..."):
        doc = Document()
        _set_normal_style(doc)

        section = doc.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        title = doc.add_heading("Meeting Summary", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph("Generated meeting notes")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_run = subtitle.runs[0]
        subtitle_run.italic = True
        subtitle_run.font.color.rgb = RGBColor(90, 90, 90)

        _add_meta_line(
            doc,
            "Date and time",
            transcript_data.get("date", datetime.now().strftime("%Y-%m-%d %H:%M")),
        )
        _add_meta_line(doc, "Source file", transcript_data.get("filename", "Unknown"))

        doc.add_heading("Summary", level=1)
        doc.add_paragraph(analysis.summary or "No summary was generated.")

        if analysis.key_decisions:
            doc.add_heading("Key Decisions", level=1)
            for decision in analysis.key_decisions:
                doc.add_paragraph(decision, style="List Bullet")

        if analysis.action_items:
            doc.add_heading("Action Items", level=1)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Task"
            hdr_cells[1].text = "Owner"
            hdr_cells[2].text = "Deadline"
            hdr_cells[3].text = "Priority"

            for item in analysis.action_items:
                row_cells = table.add_row().cells
                row_cells[0].text = item.task
                row_cells[1].text = item.owner
                row_cells[2].text = item.deadline or "Not specified"
                row_cells[3].text = item.priority or "Not specified"

        if analysis.open_questions:
            doc.add_heading("Open Questions", level=1)
            for question in analysis.open_questions:
                doc.add_paragraph(question, style="List Bullet")

        doc.add_paragraph()
        footer = doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        footer_run = footer.runs[0]
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor(110, 110, 110)

        os.makedirs("outputs", exist_ok=True)
        word_path = os.path.abspath(f"outputs/meeting_minutes_{transcript_data['meeting_id']}.docx")
        doc.save(word_path)

        st.success("Word meeting summary generated.")
        return word_path
